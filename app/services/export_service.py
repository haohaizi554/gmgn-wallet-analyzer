from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.domain.models import WalletReport
from app.services.completeness_service import CompletenessService, assert_no_empty_export_cells
from app.utils.paths import OUTPUT_DIR, REPORT_DIR, ensure_runtime_dirs
from app.utils.text import excel_cell
from app.utils.time_utils import format_datetime
from app.utils.validators import short_address
from app.version import __version__

HEADER_FILL = PatternFill("solid", fgColor="0F766E")
HEADER_FONT = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
BODY_FONT = Font(name="Microsoft YaHei", size=10)
GREEN = Font(name="Microsoft YaHei", color="15803D")
RED = Font(name="Microsoft YaHei", color="B91C1C")
ORANGE = Font(name="Microsoft YaHei", color="C2410C")
THIN = Border(
    left=Side(style="thin", color="D1D5DB"),
    right=Side(style="thin", color="D1D5DB"),
    top=Side(style="thin", color="D1D5DB"),
    bottom=Side(style="thin", color="D1D5DB"),
)


class ExportService:
    def __init__(self) -> None:
        self.completeness = CompletenessService()

    def export(self, report: WalletReport) -> tuple[Path, Path]:
        ensure_runtime_dirs()
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        short = short_address(report.request.wallet_address, 6, 4).replace("...", "")
        xlsx = OUTPUT_DIR / f"GMGN_wallet_{short}_{stamp}.xlsx"
        json_path = REPORT_DIR / f"report_{short}_{stamp}.json"
        token_rows = []
        for idx, token in enumerate(report.tokens, start=1):
            row = self.completeness.token_export_row(token)
            row["#"] = idx
            token_rows.append(row)
        platforms = {t.token_address: str(t.source_platform.export("text")) for t in report.tokens}
        trade_rows = [
            self.completeness.trade_export_row(tr, platforms.get(tr.token_address, "未知来源"))
            for tr in report.trades
        ]
        warning_rows = [
            {
                "钱包": w.wallet_address,
                "Token": w.token_symbol or "未知",
                "代币合约": w.token_address,
                "字段": w.field_name,
                "最终值": w.final_value,
                "状态": w.status,
                "原因": w.reason or "无",
                "数据来源": w.source,
                "是否估算": "是" if w.estimated else "否",
            }
            for w in report.warnings
        ]
        if not warning_rows:
            warning_rows = [
                {
                    "钱包": report.request.wallet_address,
                    "Token": "无",
                    "代币合约": "无",
                    "字段": "无",
                    "最终值": "本次无异常/补全记录",
                    "状态": "KNOWN",
                    "原因": "完整性检查未发现需要审计的 fallback",
                    "数据来源": "completeness_service",
                    "是否估算": "否",
                }
            ]
        scope_rows = [
            {
                "钱包": report.scope.wallet_address,
                "链": report.scope.chain,
                "报告周期": report.scope.period,
                "开始时间": format_datetime(report.scope.start_time) if report.scope.start_time else "全部",
                "结束时间": format_datetime(report.scope.end_time),
                "实际最早交易": format_datetime(report.scope.earliest_trade_ts) if report.scope.earliest_trade_ts else "无交易",
                "实际最晚交易": format_datetime(report.scope.latest_trade_ts) if report.scope.latest_trade_ts else "无交易",
                "报告交易数": report.scope.report_trade_count,
                "历史追溯交易数": report.scope.history_trade_count,
                "API页数": report.scope.api_pages,
                "Token数量": report.scope.token_count,
                "是否被上限截断": "是" if report.scope.truncated_by_max else "否",
                "生成时间": report.scope.generated_at or datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "版本": report.scope.version or __version__,
                "Excel路径": str(xlsx),
                "JSON路径": str(json_path),
                "主历史数据源": report.scope.primary_history_provider or "无",
                "Fallback数据源": report.scope.fallback_provider or "无",
                "请求开始": format_datetime(report.scope.start_time) if report.scope.start_time else "全部",
                "请求结束": format_datetime(report.scope.end_time) if report.scope.end_time else "全部",
                "实际覆盖开始": report.scope.actual_coverage_start or "无",
                "实际覆盖结束": report.scope.actual_coverage_end or "无",
                "Coverage Complete": report.scope.coverage_complete or "否",
                "Verified Empty": report.scope.verified_empty or "否",
                "History Pages": report.scope.history_pages,
                "History Transactions": report.scope.history_transactions,
                "Fallback Used": report.scope.fallback_used or "否",
            }
        ]
        raw_rows = [
            {
                "请求数": report.api_stats.requests,
                "缓存命中": report.api_stats.cache_hits,
                "缓存未命中": report.api_stats.cache_misses,
                "缓存命中率": report.api_stats.cache_hit_rate,
                "429重试": report.api_stats.retries_429,
                "错误数": report.api_stats.errors,
                "原始JSON目录": "; ".join(report.raw_paths) if report.raw_paths else "未保存",
                "GMGN利润周期": report.gmgn_profit.period or "接口不支持该周期",
                "GMGN利润状态": report.gmgn_profit.status.value,
                "GMGN统计周期": report.gmgn_stats.period or "GMGN 未提供",
                "任务状态": report.status.value,
                "耗时秒": round(report.elapsed_seconds, 2),
            }
        ]
        for rows in (token_rows, trade_rows, warning_rows, scope_rows, raw_rows):
            assert_no_empty_export_cells(rows)
        audit_rows = self._audit_rows(report)
        conflict_rows = self._conflict_rows(report)
        provider_rows = self._provider_rows(report)
        coverage_rows = self._coverage_rows(report)
        self._write_workbook(
            xlsx,
            [
                ("代币分析", token_rows),
                ("交易明细", trade_rows or [self._empty_trade_placeholder(report.request.wallet_address)]),
                ("异常与补全", warning_rows),
                ("原始概要", raw_rows),
                ("数据验证", audit_rows),
                ("冲突与缺失", conflict_rows),
                ("数据源统计", provider_rows),
                ("覆盖率", coverage_rows),
            ],
            hide_suffix={"代币分析": 11},
        )
        payload = self._json_payload(report, token_rows, trade_rows)
        json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
        report.excel_path = str(xlsx)
        report.json_path = str(json_path)
        return xlsx, json_path

    def export_batch(self, reports: list[WalletReport], job_id: str) -> Path:
        ensure_runtime_dirs()
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        xlsx = OUTPUT_DIR / f"GMGN_batch_{job_id[:8]}_{stamp}.xlsx"
        summary_rows = []
        token_rows = []
        trade_rows = []
        warning_rows = []
        scope_rows = []
        perf_rows = []
        for report in reports:
            s = report.summary
            summary_rows.append(
                {
                    "钱包": report.request.wallet_address,
                    "状态": report.status.value,
                    "Token数": s.token_count,
                    "买入": s.buy_count,
                    "卖出": s.sell_count,
                    "特殊获得": s.special_acquisition_count,
                    "缺失成本": s.missing_cost_count,
                    "耗时秒": round(report.elapsed_seconds, 2),
                    "Excel": report.excel_path or "见本文件",
                    "请求数": report.api_stats.requests,
                    "缓存命中率": report.api_stats.cache_hit_rate,
                }
            )
            for idx, token in enumerate(report.tokens, start=1):
                row = self.completeness.token_export_row(token)
                row["#"] = idx
                row["钱包"] = token.wallet_address
                token_rows.append(row)
            platforms = {t.token_address: str(t.source_platform.export("text")) for t in report.tokens}
            for tr in report.trades:
                trade_rows.append(self.completeness.trade_export_row(tr, platforms.get(tr.token_address, "未知来源")))
            for w in report.warnings:
                warning_rows.append(
                    {
                        "钱包": w.wallet_address,
                        "Token": w.token_symbol or "未知",
                        "代币合约": w.token_address,
                        "字段": w.field_name,
                        "最终值": w.final_value,
                        "状态": w.status,
                        "原因": w.reason or "无",
                        "数据来源": w.source,
                        "是否估算": "是" if w.estimated else "否",
                    }
                )
            scope_rows.append(
                {
                    "钱包": report.scope.wallet_address,
                    "链": report.scope.chain,
                    "报告周期": report.scope.period,
                    "开始时间": format_datetime(report.scope.start_time) if report.scope.start_time else "全部",
                    "结束时间": format_datetime(report.scope.end_time),
                    "实际最早交易": format_datetime(report.scope.earliest_trade_ts) if report.scope.earliest_trade_ts else "无交易",
                    "实际最晚交易": format_datetime(report.scope.latest_trade_ts) if report.scope.latest_trade_ts else "无交易",
                    "报告交易数": report.scope.report_trade_count,
                    "历史追溯交易数": report.scope.history_trade_count,
                    "API页数": report.scope.api_pages,
                    "Token数量": report.scope.token_count,
                    "是否被上限截断": "是" if report.scope.truncated_by_max else "否",
                    "生成时间": report.scope.generated_at or datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    "版本": report.scope.version or __version__,
                    "Excel路径": str(xlsx),
                    "JSON路径": report.json_path or "见单钱包导出",
                    "主历史数据源": report.scope.primary_history_provider or "无",
                    "Fallback数据源": report.scope.fallback_provider or "无",
                    "请求开始": format_datetime(report.scope.start_time) if report.scope.start_time else "全部",
                    "请求结束": format_datetime(report.scope.end_time) if report.scope.end_time else "全部",
                    "实际覆盖开始": report.scope.actual_coverage_start or "无",
                    "实际覆盖结束": report.scope.actual_coverage_end or "无",
                    "Coverage Complete": report.scope.coverage_complete or "否",
                    "Verified Empty": report.scope.verified_empty or "否",
                    "History Pages": report.scope.history_pages,
                    "History Transactions": report.scope.history_transactions,
                    "Fallback Used": report.scope.fallback_used or "否",
                }
            )
            perf_rows.append(
                {
                    "钱包": report.request.wallet_address,
                    "请求数": report.api_stats.requests,
                    "缓存命中": report.api_stats.cache_hits,
                    "缓存未命中": report.api_stats.cache_misses,
                    "缓存命中率": report.api_stats.cache_hit_rate,
                    "429重试": report.api_stats.retries_429,
                    "错误数": report.api_stats.errors,
                    "耗时秒": round(report.elapsed_seconds, 2),
                    "Job ID": job_id,
                    "任务状态": report.status.value,
                }
            )
        if not warning_rows:
            warning_rows = [
                {
                    "钱包": reports[0].request.wallet_address if reports else "无",
                    "Token": "无",
                    "代币合约": "无",
                    "字段": "无",
                    "最终值": "本次无异常/补全记录",
                    "状态": "KNOWN",
                    "原因": "完整性检查未发现需要审计的 fallback",
                    "数据来源": "completeness_service",
                    "是否估算": "否",
                }
            ]
        for rows in (summary_rows, token_rows, trade_rows, warning_rows, scope_rows, perf_rows):
            if rows:
                assert_no_empty_export_cells(rows)
        self._write_workbook(
            xlsx,
            [
                ("代币分析", token_rows or [{"#": 0, "币种": "无", "代币名称": "无", "代币合约": "无", "来源平台": "无", "钱包": reports[0].request.wallet_address if reports else "无"}]),
                ("交易明细", trade_rows or [self._empty_trade_placeholder(reports[0].request.wallet_address if reports else "无")]),
                ("异常补全", warning_rows),
                ("数据源统计", self._batch_provider_rows(reports)),
            ],
            hide_suffix={"代币分析": 11},
        )
        return xlsx

    def _empty_trade_placeholder(self, wallet: str) -> dict[str, Any]:
        row = {
            "时间": "无交易",
            "类型": "无",
            "币种": "无",
            "代币合约": "无",
            "数量": "无交易",
            "USD金额": "无交易",
            "SOL金额": "无交易",
            "价格USD": "无交易",
            "Gas USD": "无交易",
            "Gas SOL": "无交易",
            "单笔盈亏": "无交易",
            "来源平台": "无",
            "钱包": wallet,
            "TxHash": "无",
            "是否报告范围": "是",
        }
        assert_no_empty_export_cells([row])
        return row

    def _write_workbook(self, path: Path, sheets: list[tuple[str, list[dict[str, Any]]]], hide_suffix: dict[str, int] | None = None) -> None:
        wb = Workbook()
        default = wb.active
        wb.remove(default)
        profit_headers = {"已实现盈亏", "总盈亏", "总盈亏%", "单笔盈亏", "本地FIFO已实现"}
        for title, rows in sheets:
            ws = wb.create_sheet(title)
            headers = list(rows[0].keys()) if rows else ["空"]
            ws.append(headers)
            for cell in ws[1]:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center")
            for row in rows:
                values = [excel_cell(row.get(h, "GMGN 未提供")) for h in headers]
                ws.append(values)
                excel_row = ws.max_row
                for col, header in enumerate(headers, start=1):
                    cell = ws.cell(excel_row, col)
                    cell.font = BODY_FONT
                    cell.border = THIN
                    cell.alignment = Alignment(vertical="center")
                    value = cell.value
                    text = str(value).lstrip("'") if isinstance(value, str) else value
                    if header in profit_headers and isinstance(text, str):
                        if text.startswith("-") or "亏损" in text:
                            cell.font = RED
                        elif text not in {"未实现", "无法验证历史成本", "转入", "转出", "GMGN 未提供", "无交易", "$0.0000"} and not text.startswith("不适用"):
                            if any(ch.isdigit() for ch in text) and not text.startswith("-"):
                                cell.font = GREEN
                    if header in {"状态", "获得方式", "获得状态"} and str(value) not in {"SUCCESS", "BUY", "VERIFIED"}:
                        if str(value) in {"WARNING", "TRANSFER_IN", "BRIDGE", "WRAPPED", "FAILED"}:
                            cell.font = ORANGE
            ws.auto_filter.ref = ws.dimensions
            ws.freeze_panes = "A2"
            ws.row_dimensions[1].height = 22
            for col, header in enumerate(headers, start=1):
                max_len = len(str(header))
                for row in ws.iter_rows(min_row=2, min_col=col, max_col=col, values_only=True):
                    max_len = max(max_len, len(str(row[0])) if row[0] is not None else 0)
                width = min(max(12, max_len + 2), 48)
                if "合约" in header or "TxHash" in header or "钱包" in header:
                    width = max(width, 28)
                ws.column_dimensions[get_column_letter(col)].width = width
            hidden = (hide_suffix or {}).get(title, 0)
            if hidden and headers:
                start = max(1, len(headers) - hidden + 1)
                for col in range(start, len(headers) + 1):
                    ws.column_dimensions[get_column_letter(col)].hidden = True
        wb.save(path)

    def _audit_rows(self, report: WalletReport) -> list[dict[str, Any]]:
        rows: list[dict[str, Any]] = []
        for token in report.tokens:
            if token.audit_rows:
                rows.extend(token.audit_rows)
            else:
                rows.append(
                    {
                        "钱包": token.wallet_address,
                        "Mint": token.token_address,
                        "字段": "首买验证状态",
                        "最终值": token.first_buy_verify_status or "无",
                        "状态": token.first_buy_time.status.value,
                        "主数据源": token.first_buy_source or "无",
                        "其它数据源": "无",
                        "Source A Value": str(token.first_buy_time.export("text")),
                        "Source B Value": "无",
                        "差异": token.first_buy_time.reason or "无",
                        "证据": token.first_buy_time.source or "无",
                        "备注": token.acquisition.reason or "无",
                    }
                )
        if not rows:
            rows = [
                {
                    "钱包": report.request.wallet_address,
                    "Mint": "无",
                    "字段": "无",
                    "最终值": "无 Token",
                    "状态": "KNOWN",
                    "主数据源": "local",
                    "其它数据源": "无",
                    "Source A Value": "无",
                    "Source B Value": "无",
                    "差异": "无",
                    "证据": "无",
                    "备注": "无",
                }
            ]
        assert_no_empty_export_cells(rows)
        return rows

    def _conflict_rows(self, report: WalletReport) -> list[dict[str, Any]]:
        interesting = {"CONFLICT", "UNRESOLVED", "ESTIMATED", "NOT_APPLICABLE"}
        rows = []
        for token in report.tokens:
            fields = [
                ("入场市值", token.market_cap),
                ("创建时间", token.created_at),
                ("当前市值", token.current_market_cap),
                ("首笔买入", token.first_buy_display),
                ("来源平台", token.source_platform),
            ]
            for name, audited in fields:
                if audited.status.value in interesting or audited.estimated:
                    rows.append(
                        {
                            "钱包": token.wallet_address,
                            "Mint": token.token_address,
                            "字段": name,
                            "状态": audited.status.value,
                            "最终值": str(audited.export("text")),
                            "来源": audited.source,
                            "备注": audited.reason or "无",
                        }
                    )
        if not rows:
            rows = [
                {
                    "钱包": report.request.wallet_address,
                    "Mint": "无",
                    "字段": "无",
                    "状态": "KNOWN",
                    "最终值": "本次无冲突或缺失",
                    "来源": "completeness_service",
                    "备注": "无",
                }
            ]
        assert_no_empty_export_cells(rows)
        return rows

    def _provider_rows(self, report: WalletReport) -> list[dict[str, Any]]:
        rows = []
        for item in report.provider_metrics or []:
            rows.append(
                {
                    "数据源": item.get("provider") or "无",
                    "请求数": item.get("request_count", 0),
                    "缓存命中": item.get("cache_hit", 0),
                    "成功": item.get("success", 0),
                    "429": item.get("429", 0),
                    "5xx": item.get("5xx", 0),
                    "熔断次数": item.get("circuit_open", 0),
                    "降级次数": item.get("fallback_count", 0),
                    "耗时秒": item.get("latency_seconds", 0),
                    "估算Credits": item.get("estimated_credits", "不适用"),
                }
            )
        if not rows:
            rows = [
                {
                    "数据源": "GMGN",
                    "请求数": report.api_stats.requests,
                    "缓存命中": report.api_stats.cache_hits,
                    "成功": max(0, report.api_stats.requests - report.api_stats.errors),
                    "429": report.api_stats.retries_429,
                    "5xx": 0,
                    "熔断次数": 0,
                    "降级次数": 0,
                    "耗时秒": round(report.elapsed_seconds, 2),
                    "估算Credits": "不适用",
                }
            ]
        assert_no_empty_export_cells(rows)
        return rows

    def _batch_provider_rows(self, reports: list[WalletReport]) -> list[dict[str, Any]]:
        merged: dict[str, dict[str, Any]] = {}
        for report in reports:
            for row in self._provider_rows(report):
                name = str(row["数据源"])
                cur = merged.setdefault(
                    name,
                    {"数据源": name, "请求数": 0, "缓存命中": 0, "成功": 0, "429": 0, "5xx": 0, "熔断次数": 0, "降级次数": 0, "耗时秒": 0, "估算Credits": 0},
                )
                for key in ("请求数", "缓存命中", "成功", "429", "5xx", "熔断次数", "降级次数", "耗时秒"):
                    cur[key] = (cur.get(key) or 0) + (row.get(key) or 0)
                credit = row.get("估算Credits")
                if isinstance(credit, (int, float)):
                    cur["估算Credits"] = (cur.get("估算Credits") or 0) + credit
        rows = list(merged.values()) or [{"数据源": "无", "请求数": 0, "缓存命中": 0, "成功": 0, "429": 0, "5xx": 0, "熔断次数": 0, "降级次数": 0, "耗时秒": 0, "估算Credits": "不适用"}]
        assert_no_empty_export_cells(rows)
        return rows

    def _coverage_rows(self, report: WalletReport) -> list[dict[str, Any]]:
        cov = report.coverage or {}
        row = {
            "钱包": report.request.wallet_address,
            "字段完整率": "100%",
            "事实验证率": f"{float(cov.get('verified_rate') or 0)*100:.1f}%",
            "估算": f"{float(cov.get('estimated_rate') or 0)*100:.1f}%",
            "未解析": f"{float(cov.get('unresolved_rate') or 0)*100:.1f}%",
            "VERIFIED": cov.get("verified_fields", 0),
            "CONSENSUS": cov.get("consensus_fields", 0),
            "DIRECT": cov.get("direct_fields", 0),
            "DERIVED": cov.get("derived_fields", 0),
            "ESTIMATED": cov.get("estimated_fields", 0),
            "UNRESOLVED": cov.get("unresolved_fields", 0),
            "CONFLICT": cov.get("conflict_fields", 0),
        }
        assert_no_empty_export_cells([row])
        return [row]

    def _json_payload(self, report: WalletReport, token_rows: list[dict[str, Any]], trade_rows: list[dict[str, Any]]) -> dict[str, Any]:
        return {
            "meta": {
                "version": __version__,
                "generated_at": datetime.now().isoformat(timespec="seconds"),
                "task_id": report.task_id,
                "status": report.status.value,
                "elapsed_seconds": report.elapsed_seconds,
            },
            "wallet": {
                "address": report.request.wallet_address,
                "chain": report.request.chain,
                "period": report.request.period.value,
            },
            "summary": {
                "token_count": report.summary.token_count,
                "buy_count": report.summary.buy_count,
                "sell_count": report.summary.sell_count,
                "gmgn_realized_profit": report.summary.gmgn_realized_profit.to_audit_dict(),
                "gmgn_total_profit": report.summary.gmgn_total_profit.to_audit_dict(),
                "open_positions": report.summary.open_positions,
                "gas_total_usd": report.summary.gas_total_usd.to_audit_dict(),
                "special_acquisition_count": report.summary.special_acquisition_count,
                "missing_cost_count": report.summary.missing_cost_count,
            },
            "tokens": token_rows,
            "trades": trade_rows,
            "warnings": [
                {
                    "wallet": w.wallet_address,
                    "token": w.token_symbol,
                    "token_address": w.token_address,
                    "field": w.field_name,
                    "final_value": w.final_value,
                    "status": w.status,
                    "reason": w.reason,
                    "source": w.source,
                    "estimated": w.estimated,
                }
                for w in report.warnings
            ],
            "scope": report.scope.__dict__,
            "api_stats": {
                "requests": report.api_stats.requests,
                "cache_hits": report.api_stats.cache_hits,
                "cache_misses": report.api_stats.cache_misses,
                "cache_hit_rate": report.api_stats.cache_hit_rate,
                "retries_429": report.api_stats.retries_429,
                "errors": report.api_stats.errors,
            },
        }
